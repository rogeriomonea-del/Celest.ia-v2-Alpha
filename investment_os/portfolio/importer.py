"""Importação segura de carteira (XLSX/CSV) — Fase 5.

Pipeline: validar arquivo -> processar em diretório temporário -> remover PII
(determinístico, antes de qualquer persistência) -> extrair posições ->
normalizar -> resolver instrumentos (FCA/cadastro CVM) -> prévia -> correção ->
confirmação explícita -> snapshot imutável.

Garantias:
- O arquivo bruto NUNCA vai para LLM, logs, repositório ou telemetria; após o
  parse ele é removido e resta apenas o sha256 + metadados sanitizados.
- Ticker não resolvido nunca é aproximado silenciosamente.
- Custo ausente = 'desconhecido' (NULL), nunca zero.
- Reimportação idêntica é idempotente (hash do conteúdo normalizado).
- Falha parcial exige aceite explícito (accept_partial) para confirmar.
- Sem layout oficial real da B3 disponível, os adaptadores são versionados e
  honestos: formato não reconhecido gera erro claro, não um palpite.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from . import db as pdb
from .pii import scrub_value

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
_TICKER_RE = re.compile(r"^[A-Z]{4}\d{1,2}[A-Z]?$")

# Assinaturas de arquivo (magic bytes)
_ZIP_MAGIC = b"PK\x03\x04"
_EXE_MAGICS = (b"MZ", b"\x7fELF")


class ImportError_(Exception):
    """Erro de importação com mensagem segura (sem conteúdo do arquivo)."""


@dataclass
class ParsedRow:
    row_index: int
    ticker_raw: str
    quantity: float | None
    avg_cost: float | None
    currency: str
    data_base: str | None
    issues: list[str]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def validate_file(path: Path, declared_name: str) -> str:
    """Valida extensão x assinatura, tamanho e tipo. Retorna kind: xlsx|csv."""
    size = path.stat().st_size
    if size == 0:
        raise ImportError_("arquivo vazio")
    if size > MAX_UPLOAD_BYTES:
        raise ImportError_(f"arquivo excede o limite de {MAX_UPLOAD_BYTES // (1024 * 1024)}MB")
    head = path.read_bytes()[:8]
    ext = declared_name.lower().rsplit(".", 1)[-1] if "." in declared_name else ""
    if any(head.startswith(m) for m in _EXE_MAGICS):
        raise ImportError_("arquivo executável rejeitado")
    if ext == "xlsx":
        if not head.startswith(_ZIP_MAGIC):
            raise ImportError_("extensão .xlsx com assinatura incompatível (MIME falso)")
        import zipfile

        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
        except zipfile.BadZipFile as exc:
            raise ImportError_("xlsx corrompido") from exc
        if "[Content_Types].xml" not in names:
            raise ImportError_("xlsx inválido (sem Content_Types)")
        if any("vbaProject" in n for n in names):
            raise ImportError_("planilha com macros rejeitada")
        return "xlsx"
    if ext == "csv":
        if head.startswith(_ZIP_MAGIC):
            raise ImportError_("extensão .csv com assinatura de zip (MIME falso)")
        try:
            path.read_bytes().decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                path.read_bytes().decode("latin-1")
            except UnicodeDecodeError as exc:
                raise ImportError_("csv com encoding não reconhecido") from exc
        return "csv"
    raise ImportError_(
        f"formato .{ext or '?'} não suportado (aceitos: xlsx, csv; PDF fora do MVP — ADR-0004)"
    )


def _clean_number(raw) -> float | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip().replace("R$", "").replace(" ", "")
    if not s or s in {"-", "—"}:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _cell_is_formula(value) -> bool:
    return isinstance(value, str) and value.strip().startswith(("=", "+@", "@"))


def _norm_header(h) -> str:
    import unicodedata

    s = str(h or "").strip().lower()
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


# Mapeamento de colunas por adaptador. b3_consolidado_v1 reconhece cabeçalhos
# usuais do Relatório Consolidado da Área do Investidor B3 (posição em renda
# variável); é best-effort versionado — sem match, o erro é honesto.
_ADAPTERS: dict[str, dict[str, list[str]]] = {
    "b3_consolidado_v1": {
        "ticker": ["codigo de negociacao", "codigo negociacao"],
        "quantity": ["quantidade", "quantidade disponivel", "qtde."],
        "avg_cost": ["preco medio", "preco medio (r$)", "preco medio ponderado"],
        "data_base": ["data-base", "data base", "data referencia"],
    },
    "generic_csv_v1": {
        "ticker": ["ticker", "ativo", "codigo"],
        "quantity": ["quantidade", "qtd", "qty"],
        "avg_cost": ["preco_medio", "custo_medio", "avg_cost"],
        "data_base": ["data_base", "data-base"],
    },
}


def _match_adapter(headers: list[str]) -> tuple[str, dict[str, int]] | None:
    norm = [_norm_header(h) for h in headers]
    for adapter, spec in _ADAPTERS.items():
        cols: dict[str, int] = {}
        for field_name, aliases in spec.items():
            for alias in aliases:
                if alias in norm:
                    cols[field_name] = norm.index(alias)
                    break
        if "ticker" in cols and "quantity" in cols:
            return adapter, cols
    return None


def _extract_rows(path: Path, kind: str) -> tuple[str, list[list]]:
    """Extrai matriz de células. XLSX: primeira aba com adaptador reconhecido."""
    if kind == "csv":
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        delimiter = ";" if text.splitlines()[0].count(";") >= text.splitlines()[0].count(",") else ","
        return "csv", list(csv.reader(io.StringIO(text), delimiter=delimiter))
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        for ws in wb.worksheets:
            rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
            for header_idx in range(min(5, len(rows))):
                if _match_adapter([str(v) for v in rows[header_idx]]):
                    return ws.title, rows[header_idx:]
        return wb.sheetnames[0] if wb.sheetnames else "?", []
    finally:
        wb.close()


def parse_file(path: Path, kind: str) -> tuple[str, list[ParsedRow], int]:
    """Parse + sanitização. Retorna (adapter, linhas, pii_removida)."""
    _sheet, rows = _extract_rows(path, kind)
    if not rows:
        raise ImportError_("nenhuma tabela de posições reconhecida no arquivo")
    match = _match_adapter([str(v) for v in rows[0]])
    if match is None:
        raise ImportError_(
            "layout não reconhecido pelos adaptadores versionados "
            f"({', '.join(_ADAPTERS)}); cabeçalhos esperados: ticker/código de "
            "negociação + quantidade"
        )
    adapter, cols = match
    parsed: list[ParsedRow] = []
    pii_total = 0
    today = date.today().isoformat()
    for i, row in enumerate(rows[1:], start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        issues: list[str] = []
        if any(_cell_is_formula(v) for v in row):
            parsed.append(ParsedRow(i, "", None, None, "BRL", None, ["fórmula em célula rejeitada (segurança)"]))
            continue

        def cell(field_name: str):
            idx = cols.get(field_name)
            return row[idx] if idx is not None and idx < len(row) else None

        ticker_raw, n1 = scrub_value(cell("ticker"))
        pii_total += n1
        ticker_raw = ticker_raw.strip().upper()
        quantity = _clean_number(cell("quantity"))
        avg_cost = _clean_number(cell("avg_cost"))
        data_base_raw, n2 = scrub_value(cell("data_base"))
        pii_total += n2
        data_base = _parse_date(data_base_raw)

        if not ticker_raw:
            issues.append("ticker ausente")
        if quantity is None:
            issues.append("quantidade ausente ou inválida")
        elif quantity < 0:
            issues.append("quantidade negativa exige revisão")
        if avg_cost is not None and avg_cost < 0:
            issues.append("custo negativo inválido")
        if data_base and data_base > today:
            issues.append("data-base no futuro")
        parsed.append(ParsedRow(i, ticker_raw, quantity, avg_cost, "BRL", data_base, issues))
    return adapter, parsed, pii_total


def _parse_date(s: str) -> str | None:
    s = (s or "").strip()
    if not s:
        return None
    m = re.match(r"^(\d{2})/(\d{2})/(\d{4})$", s)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    m = re.match(r"^\d{4}-\d{2}-\d{2}", s)
    return m.group(0) if m else None


# ---------------------------------------------------------------- resolução


def load_reference() -> tuple[dict, dict]:
    """Listagens FCA (ticker -> cnpj/classe/nome) e cadastro (cnpj -> cd_cvm/setor)."""
    from .. import config
    import pandas as pd

    listings_path = config.SILVER_DIR / "listings.parquet"
    by_ticker: dict[str, dict] = {}
    if listings_path.exists():
        for _, r in pd.read_parquet(listings_path).iterrows():
            by_ticker[str(r["ticker"]).upper()] = {
                "cnpj": r["cnpj"], "classe": r["classe"], "nome": r["nome"],
            }
    cad_path = config.BRONZE_DIR / "cvm_dados_abertos" / "cad_cia_aberta.csv"
    by_cnpj: dict[str, dict] = {}
    if cad_path.exists():
        cad = pd.read_csv(cad_path, sep=";", encoding="latin-1", dtype=str)
        for _, r in cad.iterrows():
            by_cnpj[str(r["CNPJ_CIA"])] = {
                "cd_cvm": str(r["CD_CVM"]).lstrip("0"), "setor": r["SETOR_ATIV"],
            }
    return by_ticker, by_cnpj


def resolve_instrument(ticker: str, by_ticker: dict, by_cnpj: dict) -> dict:
    """Resolução determinística SEM aproximação silenciosa."""
    if not _TICKER_RE.match(ticker or ""):
        return {"status": "unknown", "confidence": "BAIXA", "method": "unresolved",
                "reason": f"'{ticker}' não tem formato de código de negociação B3"}
    hit = by_ticker.get(ticker)
    if hit:
        extra = by_cnpj.get(hit["cnpj"], {})
        return {
            "status": "ok", "confidence": "ALTA", "method": "fca_listing",
            "ticker": ticker, "cnpj": hit["cnpj"], "classe": hit["classe"],
            "cd_cvm": extra.get("cd_cvm"), "setor": extra.get("setor"),
            "asset_class": "acao_br",
        }
    if ticker.endswith("11"):
        return {"status": "ambiguous", "confidence": "BAIXA", "method": "unresolved",
                "reason": "final 11 pode ser FII, unit ou ETF — sem cadastro oficial "
                          "integrado para confirmar; requer confirmação manual",
                "asset_class_hint": "fii_ou_unit_ou_etf"}
    return {"status": "unknown", "confidence": "BAIXA", "method": "unresolved",
            "reason": "ticker não encontrado nas listagens oficiais (FCA/CVM) ingeridas"}


# ---------------------------------------------------------------- pipeline


def start_import(conn: sqlite3.Connection, portfolio_id: int, file_path: Path,
                 declared_name: str, *, keep_source_file: bool = False) -> dict:
    """Valida, sanitiza, resolve e grava a prévia. Remove o arquivo temporário."""
    kind = validate_file(file_path, declared_name)
    digest = sha256_file(file_path)
    size = file_path.stat().st_size
    try:
        adapter, rows, pii_count = parse_file(file_path, kind)
    except ImportError_ as exc:
        cur = conn.execute(
            "INSERT INTO portfolio_import (portfolio_id, created_at, file_sha256, file_bytes,"
            " file_kind, adapter, status, pii_removed_count, error) VALUES (?,?,?,?,?,?,?,?,?)",
            (portfolio_id, pdb.utcnow(), digest, size, kind, "?", "failed", 0, str(exc)),
        )
        conn.commit()
        pdb.audit(conn, "import_failed", import_id=cur.lastrowid, sha256=digest, error=str(exc))
        raise
    finally:
        if not keep_source_file:
            file_path.unlink(missing_ok=True)  # política de retenção: só hash

    by_ticker, by_cnpj = load_reference()
    data_bases = {r.data_base for r in rows if r.data_base}
    cur = conn.execute(
        "INSERT INTO portfolio_import (portfolio_id, created_at, file_sha256, file_bytes,"
        " file_kind, adapter, status, data_base, pii_removed_count) VALUES (?,?,?,?,?,?,?,?,?)",
        (portfolio_id, pdb.utcnow(), digest, size, kind, adapter, "previewed",
         max(data_bases) if data_bases else None, pii_count),
    )
    import_id = cur.lastrowid

    seen: set[tuple[str, str | None]] = set()
    for r in rows:
        if r.issues:
            status, reason, resolution = "rejected", "; ".join(r.issues), {}
        else:
            resolution = resolve_instrument(r.ticker_raw, by_ticker, by_cnpj)
            status = {"ok": "ok", "ambiguous": "ambiguous", "unknown": "unknown"}[resolution["status"]]
            reason = resolution.get("reason", "")
            key = (r.ticker_raw, r.data_base)
            if key in seen:
                status, reason = "duplicate", "linha duplicada (mesmo ticker e data-base)"
            seen.add(key)
        parsed_payload = {
            "ticker": r.ticker_raw, "quantity": r.quantity,
            "avg_cost": r.avg_cost,
            "cost_status": "conhecido" if r.avg_cost is not None else "desconhecido",
            "currency": r.currency, "data_base": r.data_base,
            "resolution": resolution,
        }
        rcur = conn.execute(
            "INSERT INTO portfolio_import_row (import_id, row_index, parsed_json, status,"
            " reason, resolution_confidence) VALUES (?,?,?,?,?,?)",
            (import_id, r.row_index, json.dumps(parsed_payload, ensure_ascii=False),
             status, reason, resolution.get("confidence", "BAIXA")),
        )
        if resolution.get("status") == "ok":
            conn.execute(
                "INSERT INTO instrument_resolution (import_row_id, ticker, cd_cvm, cnpj,"
                " classe, asset_class, method, confidence, resolved_by)"
                " VALUES (?,?,?,?,?,?,?,?,?)",
                (rcur.lastrowid, resolution.get("ticker"), resolution.get("cd_cvm"),
                 resolution.get("cnpj"), resolution.get("classe"),
                 resolution.get("asset_class"), resolution["method"],
                 resolution["confidence"], "system"),
            )
    conn.commit()
    pdb.audit(conn, "import_previewed", import_id=import_id, sha256=digest,
              rows=len(rows), pii_removed=pii_count)
    return preview(conn, import_id)


def preview(conn: sqlite3.Connection, import_id: int) -> dict:
    imp = conn.execute("SELECT * FROM portfolio_import WHERE id=?", (import_id,)).fetchone()
    if imp is None:
        raise ImportError_("importação não encontrada")
    rows = conn.execute(
        "SELECT * FROM portfolio_import_row WHERE import_id=? ORDER BY row_index", (import_id,)
    ).fetchall()
    counts: dict[str, int] = {}
    out_rows = []
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
        out_rows.append(
            {"row_id": r["id"], "row_index": r["row_index"], "status": r["status"],
             "reason": r["reason"], "confidence": r["resolution_confidence"],
             **json.loads(r["parsed_json"])}
        )
    return {
        "import_id": import_id, "status": imp["status"], "adapter": imp["adapter"],
        "file_sha256": imp["file_sha256"], "data_base": imp["data_base"],
        "pii_removed_count": imp["pii_removed_count"], "counts": counts, "rows": out_rows,
        "requires_user_confirmation": True,
    }


def correct_row(conn: sqlite3.Connection, import_id: int, row_id: int, ticker: str,
                asset_class: str | None = None) -> dict:
    """Correção manual de mapeamento antes da confirmação."""
    row = conn.execute(
        "SELECT * FROM portfolio_import_row WHERE id=? AND import_id=?", (row_id, import_id)
    ).fetchone()
    if row is None:
        raise ImportError_("linha não encontrada")
    by_ticker, by_cnpj = load_reference()
    ticker = ticker.strip().upper()
    resolution = resolve_instrument(ticker, by_ticker, by_cnpj)
    if resolution["status"] != "ok" and asset_class:
        # usuário assume a classificação: registrada como user_override
        resolution = {"status": "ok", "confidence": "MEDIA", "method": "user_override",
                      "ticker": ticker, "asset_class": asset_class,
                      "cnpj": None, "classe": None, "cd_cvm": None, "setor": None}
    payload = json.loads(row["parsed_json"])
    payload["ticker"] = ticker
    payload["resolution"] = resolution
    status = "ok" if resolution["status"] == "ok" else resolution["status"]
    conn.execute(
        "UPDATE portfolio_import_row SET parsed_json=?, status=?, reason=?, resolution_confidence=?"
        " WHERE id=?",
        (json.dumps(payload, ensure_ascii=False), status, resolution.get("reason", ""),
         resolution.get("confidence", "BAIXA"), row_id),
    )
    if resolution["status"] == "ok":
        conn.execute(
            "INSERT INTO instrument_resolution (import_row_id, ticker, cd_cvm, cnpj, classe,"
            " asset_class, method, confidence, resolved_by) VALUES (?,?,?,?,?,?,?,?,?)",
            (row_id, ticker, resolution.get("cd_cvm"), resolution.get("cnpj"),
             resolution.get("classe"), resolution.get("asset_class"),
             resolution["method"], resolution["confidence"],
             "user" if resolution["method"] == "user_override" else "system"),
        )
    conn.commit()
    pdb.audit(conn, "import_row_corrected", import_id=import_id, row_id=row_id)
    return preview(conn, import_id)


def confirm_import(conn: sqlite3.Connection, import_id: int, *, accept_partial: bool = False) -> dict:
    """Confirmação explícita -> snapshot imutável. Idempotente por conteúdo."""
    prev = preview(conn, import_id)
    if prev["status"] == "confirmed":
        snap = conn.execute(
            "SELECT * FROM portfolio_snapshot WHERE import_id=?", (import_id,)
        ).fetchone()
        return {"snapshot_id": snap["id"], "version": snap["version"], "idempotent": True}
    ok_rows = [r for r in prev["rows"] if r["status"] == "ok"]
    problem_rows = [r for r in prev["rows"] if r["status"] not in ("ok", "duplicate")]
    if not ok_rows:
        raise ImportError_("nenhuma linha válida para confirmar")
    if problem_rows and not accept_partial:
        raise ImportError_(
            f"importação parcial: {len(problem_rows)} linha(s) com problema "
            "(ambíguas/desconhecidas/rejeitadas). Corrija-as ou confirme "
            "explicitamente com accept_partial=true — nunca silenciosamente."
        )
    imp = conn.execute("SELECT * FROM portfolio_import WHERE id=?", (import_id,)).fetchone()
    portfolio_id = imp["portfolio_id"]

    normalized = sorted(
        (r["ticker"], r["quantity"], r["avg_cost"], r["data_base"]) for r in ok_rows
    )
    content_sha = hashlib.sha256(json.dumps(normalized, default=str).encode()).hexdigest()
    existing = conn.execute(
        "SELECT * FROM portfolio_snapshot WHERE portfolio_id=? AND content_sha256=?",
        (portfolio_id, content_sha),
    ).fetchone()
    if existing:
        conn.execute("UPDATE portfolio_import SET status='confirmed' WHERE id=?", (import_id,))
        conn.commit()
        pdb.audit(conn, "import_confirmed_idempotent", import_id=import_id,
                  snapshot_id=existing["id"])
        return {"snapshot_id": existing["id"], "version": existing["version"], "idempotent": True}

    version = (conn.execute(
        "SELECT COALESCE(MAX(version), 0) AS v FROM portfolio_snapshot WHERE portfolio_id=?",
        (portfolio_id,),
    ).fetchone()["v"]) + 1
    cur = conn.execute(
        "INSERT INTO portfolio_snapshot (portfolio_id, import_id, version, created_at,"
        " data_base, content_sha256) VALUES (?,?,?,?,?,?)",
        (portfolio_id, import_id, version, pdb.utcnow(), prev["data_base"], content_sha),
    )
    snapshot_id = cur.lastrowid
    for r in ok_rows:
        res = r.get("resolution", {})
        pcur = conn.execute(
            "INSERT INTO position (snapshot_id, ticker, cd_cvm, cnpj, isin, classe,"
            " asset_class, quantity, avg_cost, cost_status, currency, source)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
            (snapshot_id, r["ticker"], res.get("cd_cvm"), res.get("cnpj"), None,
             res.get("classe"), res.get("asset_class") or "outro", r["quantity"],
             r["avg_cost"], r["cost_status"], r["currency"], f"import:{import_id}"),
        )
        conn.execute(
            "INSERT INTO position_source (position_id, import_row_id, kind) VALUES (?,?,?)",
            (pcur.lastrowid, r["row_id"], "import_row"),
        )
    conn.execute("UPDATE portfolio_import SET status='confirmed' WHERE id=?", (import_id,))
    if problem_rows:
        pdb.add_issue(conn, "import", import_id, "alta",
                      f"{len(problem_rows)} linha(s) excluída(s) do snapshot com aceite explícito do usuário")
    conn.commit()
    pdb.audit(conn, "snapshot_created", snapshot_id=snapshot_id, version=version,
              positions=len(ok_rows), content_sha256=content_sha)
    return {"snapshot_id": snapshot_id, "version": version, "idempotent": False,
            "excluded_rows": len(problem_rows)}
